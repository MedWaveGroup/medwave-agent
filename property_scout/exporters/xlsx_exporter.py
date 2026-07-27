"""XLSX exporter - one row per listing, plus a summary tab."""
from __future__ import annotations

import statistics
from pathlib import Path
from typing import List

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from ..models import Listing
from ..pipeline import RunResult

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
SECTION_FILL = PatternFill("solid", fgColor="D9E1F2")

COLUMNS = [
    ("Section", lambda l, sec: sec),
    ("Source", lambda l, sec: l.source),
    ("Agency", lambda l, sec: l.agency_name),
    ("Also listed by", lambda l, sec: l.also_listed_by),
    ("Agent", lambda l, sec: l.agent_name or ""),
    ("Phone", lambda l, sec: l.agent_phone or ""),
    ("Email", lambda l, sec: l.agent_email or ""),
    ("Title", lambda l, sec: l.title),
    ("Type", lambda l, sec: l.property_type),
    ("Address", lambda l, sec: l.address_raw),
    ("Suburb", lambda l, sec: l.suburb),
    ("City", lambda l, sec: l.city),
    ("Province", lambda l, sec: l.province),
    ("Rent (R)", lambda l, sec: l.monthly_rent),
    ("VAT basis", lambda l, sec: l.rent_basis),
    ("Additional costs", lambda l, sec: l.additional_costs),
    ("Size (m²)", lambda l, sec: l.size_sqm),
    ("Size confidence", lambda l, sec: l.size_confidence),
    ("Days on market", lambda l, sec: l.days_on_market),
    ("Age confidence", lambda l, sec: l.age_confidence),
    ("Availability", lambda l, sec: l.availability_status),
    ("Verified live", lambda l, sec: l.last_seen.strftime("%Y-%m-%d") if l.last_seen else ""),
    ("Date listed", lambda l, sec: l.date_listed.isoformat() if l.date_listed else ""),
    ("First seen", lambda l, sec: l.date_first_seen.isoformat() if l.date_first_seen else ""),
    ("Flags", lambda l, sec: ";".join(l.flags)),
    ("URL", lambda l, sec: l.url),
]


def _style_header(ws, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"


def _autosize(ws, max_width: int = 55) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        length = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[letter].width = min(max(12, length + 2), max_width)


def export_xlsx(result: RunResult, out_dir: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Listings"

    headers = [h for h, _ in COLUMNS]
    ws.append(headers)
    _style_header(ws, len(headers))

    sections = [
        ("Qualifying", result.main),
        ("Age Unconfirmed", result.age_unconfirmed),
        ("Needs Verification", result.needs_verification),
    ]
    all_rows: List[Listing] = []
    for sec_name, listings in sections:
        for l in listings:
            ws.append([fn(l, sec_name) for _, fn in COLUMNS])
            all_rows.append(l)
            # Hyperlink the URL cell.
            url_cell = ws.cell(row=ws.max_row, column=len(COLUMNS))
            if l.url:
                url_cell.hyperlink = l.url
                url_cell.font = Font(color="0563C1", underline="single")

    _autosize(ws)

    # --- Dropped-since-last tab -----------------------------------------
    if result.dropped_since_last:
        dws = wb.create_sheet("Dropped Since Last")
        dws.append(["Title", "Agency", "Suburb", "Reason", "URL"])
        _style_header(dws, 5)
        for d in result.dropped_since_last:
            dws.append([d.get("title", ""), d.get("agency_name", ""),
                        d.get("suburb", ""), d.get("reason", ""), d.get("url", "")])
        _autosize(dws)

    # --- Summary tab -----------------------------------------------------
    sws = wb.create_sheet("Summary")
    _write_summary(sws, result, all_rows)

    out_dir.mkdir(parents=True, exist_ok=True)
    safe_area = result.area.split(",")[0].strip().lower().replace(" ", "_")
    path = out_dir / f"property_scout_{safe_area}_{result.run_id}.xlsx"
    wb.save(path)
    return path


def _write_summary(ws, result: RunResult, listings: List[Listing]) -> None:
    ws.append(["Property Scout - Summary"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["Area", result.area])
    ws.append(["Run ID", result.run_id])
    ws.append(["Generated", result.started_at.strftime("%Y-%m-%d %H:%M")])
    ws.append(["Qualifying", len(result.main)])
    ws.append(["Age unconfirmed", len(result.age_unconfirmed)])
    ws.append(["Needs verification", len(result.needs_verification)])
    ws.append(["Dropped since last run", len(result.dropped_since_last)])
    ws.append([])

    # By suburb: count + median rent.
    ws.append(["Count & median rent by suburb"])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    ws.append(["Suburb", "Count", "Median rent (R)"])
    by_suburb: dict = {}
    for l in listings:
        by_suburb.setdefault(l.suburb or "(unknown)", []).append(l.monthly_rent)
    for suburb, rents in sorted(by_suburb.items()):
        vals = [r for r in rents if r]
        med = int(statistics.median(vals)) if vals else ""
        ws.append([suburb, len(rents), med])
    ws.append([])

    # By agency.
    ws.append(["Count by agency"])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    ws.append(["Agency", "Count"])
    by_agency: dict = {}
    for l in listings:
        by_agency[l.agency_name or "(unknown)"] = by_agency.get(l.agency_name or "(unknown)", 0) + 1
    for agency, count in sorted(by_agency.items(), key=lambda x: -x[1]):
        ws.append([agency, count])
    ws.append([])

    # By property type.
    ws.append(["Count by property type"])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    ws.append(["Type", "Count"])
    by_type: dict = {}
    for l in listings:
        by_type[l.property_type] = by_type.get(l.property_type, 0) + 1
    for ptype, count in sorted(by_type.items(), key=lambda x: -x[1]):
        ws.append([ptype, count])

    _autosize(ws)
