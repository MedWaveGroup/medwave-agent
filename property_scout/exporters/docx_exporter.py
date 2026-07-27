"""DOCX call sheet - the primary deliverable the team works from.

Layout (spec):
  * cover block with filters + source status;
  * body grouped BY AGENCY (one call covers several properties);
  * per-agency property table sorted freshest-first, with blank Called/Outcome/
    Viewing-booked columns;
  * "Age Unconfirmed", "Needs Verification" and "Dropped Since Last Run"
    sections;
  * clean, printable, portrait A4.
"""
from __future__ import annotations

from collections import defaultdict
from pathlib import Path
from typing import List, Optional

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Cm

from ..models import Listing
from ..pipeline import RunResult

NAVY = RGBColor(0x1F, 0x4E, 0x78)
GREY = RGBColor(0x55, 0x55, 0x55)


def _dom_key(l: Listing) -> int:
    return l.days_on_market if l.days_on_market is not None else 10_000


def _set_a4_portrait(section) -> None:
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    for attr in ("left_margin", "right_margin"):
        setattr(section, attr, Cm(1.5))
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)


def _heading(doc, text: str, size: int = 14, color=NAVY, space_before: int = 8):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    run.font.color.rgb = color
    return p


def _rent_str(l: Listing) -> str:
    if l.monthly_rent is None:
        return "POA"
    basis = {"excl_vat": " excl VAT", "incl_vat": " incl VAT"}.get(l.rent_basis, "")
    s = f"R{l.monthly_rent:,}{basis}"
    if "vat_pushes_over_limit" in l.flags:
        s += " *"
    return s


def _size_str(l: Listing) -> str:
    if l.size_sqm is None:
        return "?"
    mark = "" if l.size_confidence == "stated" else "~"
    return f"{mark}{l.size_sqm:g}"


def _verified_str(l: Listing) -> str:
    return l.last_seen.strftime("%d %b") if l.last_seen else "-"


def _property_table(doc, listings: List[Listing], include_blanks: bool = True) -> None:
    headers = ["Address / Suburb", "Size m²", "Rent", "Type", "DoM", "Live",
               "Called", "Outcome", "Booked"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        run = hdr[i].paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(8)

    for l in sorted(listings, key=_dom_key):
        cells = table.add_row().cells
        addr = l.address_raw or l.title or "(see listing)"
        loc = f"{addr}"
        if l.suburb and l.suburb.lower() not in addr.lower():
            loc += f"\n{l.suburb}"
        dom = str(l.days_on_market) if l.days_on_market is not None else "?"
        values = [loc, _size_str(l), _rent_str(l), l.property_type, dom,
                  _verified_str(l), "", "", ""]
        for i, v in enumerate(values):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(v))
            run.font.size = Pt(8)
        # Link the address cell to the listing.
        if l.url:
            _hyperlink_cell(cells[0], l.url)

    # Set narrow, printable widths.
    widths = [Cm(4.6), Cm(1.3), Cm(2.4), Cm(1.9), Cm(1.0), Cm(1.2),
              Cm(1.6), Cm(2.6), Cm(1.4)]
    for row in table.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = w


def _hyperlink_cell(cell, url: str) -> None:
    """Add a small clickable 'view' link under a cell's text."""
    p = cell.add_paragraph()
    from docx.oxml.shared import OxmlElement, qn
    part = cell.part
    r_id = part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "14")
    rpr.append(color)
    rpr.append(u)
    rpr.append(sz)
    new_run.append(rpr)
    t = OxmlElement("w:t")
    t.text = "view listing"
    new_run.append(t)
    hyperlink.append(new_run)
    p._p.append(hyperlink)


def export_docx(result: RunResult, out_dir: Path) -> Path:
    doc = Document()
    _set_a4_portrait(doc.sections[0])

    # Base font.
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)

    # --- cover block -----------------------------------------------------
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = title.add_run("Property Scout - Call Sheet")
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = NAVY

    sub = doc.add_paragraph()
    rs = sub.add_run(f"{result.area}   •   {result.started_at.strftime('%d %B %Y')}")
    rs.font.size = Pt(12)
    rs.font.color.rgb = GREY

    f = result.filters
    filt = doc.add_paragraph()
    filt.add_run("Filters applied: ").bold = True
    filt.add_run(
        f"rent ≤ R{f['max_rent']:,}, size ≥ {f['min_size_sqm']} m², "
        f"on market < {f['max_listing_age_days']} days, "
        f"radius {f.get('radius_km') or 'n/a'} km, "
        f"residential {'included' if f['include_residential'] else 'excluded'}."
    )

    totals = doc.add_paragraph()
    totals.add_run("Results: ").bold = True
    totals.add_run(
        f"{len(result.main)} qualifying • {len(result.age_unconfirmed)} age-unconfirmed"
        f" • {len(result.needs_verification)} need verification"
        f" • {len(result.dropped_since_last)} dropped since last run."
    )

    src = doc.add_paragraph()
    src.add_run("Sources queried: ").bold = True
    src.add_run(", ".join(result.sources_queried) or "none")
    blocked = result.sources_blocked
    if blocked:
        b = doc.add_paragraph()
        b.add_run("Sources blocked / skipped: ").bold = True
        rb = b.add_run(", ".join(blocked))
        rb.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)

    note = doc.add_paragraph()
    rn = note.add_run(
        "Rent marked * means excl-VAT rent whose VAT-inclusive cost exceeds the "
        "limit. Size prefixed ~ was parsed from description text. "
        "DoM = days on market; Live = date last confirmed on the source site."
    )
    rn.italic = True
    rn.font.size = Pt(8)
    rn.font.color.rgb = GREY

    # --- qualifying, grouped by agency ----------------------------------
    _heading(doc, "Qualifying listings - grouped by agency", size=15)
    if not result.main:
        doc.add_paragraph("No qualifying listings this run.")
    else:
        by_agency = defaultdict(list)
        for l in result.main:
            by_agency[l.agency_name or "(agency unknown)"].append(l)
        # Agencies sorted by freshest listing first.
        for agency in sorted(by_agency, key=lambda a: min(_dom_key(x) for x in by_agency[a])):
            listings = by_agency[agency]
            _heading(doc, agency, size=12, space_before=10)
            # Contact line - take richest contact among this agency's listings.
            contact = _best_contact(listings)
            cp = doc.add_paragraph()
            cp.paragraph_format.space_after = Pt(2)
            cr = cp.add_run(contact)
            cr.font.size = Pt(9)
            cr.font.color.rgb = GREY
            _property_table(doc, listings)

    # --- Age Unconfirmed -------------------------------------------------
    _heading(doc, "Age Unconfirmed", size=14, space_before=14)
    p = doc.add_paragraph()
    pr = p.add_run("Passed every other filter, but we could not establish how "
                   "long these have been listed. Ask the agent how long this has "
                   "been on the market.")
    pr.italic = True
    pr.font.size = Pt(9)
    if result.age_unconfirmed:
        _property_table(doc, result.age_unconfirmed)
    else:
        doc.add_paragraph("None.")

    # --- Needs Verification ---------------------------------------------
    _heading(doc, "Needs Verification", size=14, space_before=14)
    p = doc.add_paragraph()
    pr = p.add_run("Matched area and type but had unclear size or price "
                   "(e.g. 'Price on Application'). Confirm details with the agent.")
    pr.italic = True
    pr.font.size = Pt(9)
    if result.needs_verification:
        _property_table(doc, result.needs_verification)
    else:
        doc.add_paragraph("None.")

    # --- Dropped Since Last Run -----------------------------------------
    _heading(doc, "Dropped Since Last Run", size=14, space_before=14)
    p = doc.add_paragraph()
    pr = p.add_run("Previously exported listings that have since been let or "
                   "withdrawn - stop chasing these.")
    pr.italic = True
    pr.font.size = Pt(9)
    if result.dropped_since_last:
        t = doc.add_table(rows=1, cols=4)
        t.style = "Light List Accent 1"
        for i, h in enumerate(["Title", "Agency", "Suburb", "Reason"]):
            run = t.rows[0].cells[i].paragraphs[0].add_run(h)
            run.bold = True
            run.font.size = Pt(8)
        for d in result.dropped_since_last:
            cells = t.add_row().cells
            for i, v in enumerate([d.get("title", ""), d.get("agency_name", ""),
                                   d.get("suburb", ""), d.get("reason", "")]):
                cells[i].paragraphs[0].add_run(str(v)).font.size = Pt(8)
    else:
        doc.add_paragraph("None.")

    out_dir.mkdir(parents=True, exist_ok=True)
    safe_area = result.area.split(",")[0].strip().lower().replace(" ", "_")
    path = out_dir / f"call_sheet_{safe_area}_{result.run_id}.docx"
    doc.save(path)
    return path


def _best_contact(listings: List[Listing]) -> str:
    def score(l: Listing) -> int:
        return (3 if l.agent_phone else 0) + (2 if l.agent_email else 0) + (1 if l.agent_name else 0)
    best = max(listings, key=score)
    parts = []
    if best.agent_name:
        parts.append(best.agent_name)
    if best.agent_phone:
        parts.append(f"☎ {best.agent_phone}")
    if best.agent_email:
        parts.append(f"✉ {best.agent_email}")
    also = {l.also_listed_by for l in listings if l.also_listed_by}
    line = "  |  ".join(parts) if parts else "No direct contact captured - see listing."
    return line
